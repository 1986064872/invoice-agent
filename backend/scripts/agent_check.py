"""隔离端到端检查：只使用公开合成票据和临时数据库，不访问用户票据。"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import json
from unittest.mock import patch as mock_patch
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
scratch = tempfile.TemporaryDirectory(prefix='invoice-agent-check-')
temp = Path(scratch.name)
os.environ['APP_DATA_DIR'] = str(temp / 'data')
os.environ['APP_DB_PATH'] = str(temp / 'data' / 'test.db')
os.environ['OCR_PROVIDER'] = 'local'

from fastapi.testclient import TestClient
from app.main import app
from app.routers import agent
from app.database import SessionLocal, engine
from app.models import Invoice, LedgerEntry


@contextmanager
def mock_remote(deepseek, responses):
    with mock_patch.object(deepseek.httpx, 'Client') as factory:
        remote = factory.return_value.__enter__.return_value.post
        remote.side_effect = responses
        yield remote


def main():
    agent.SOURCE_ROOT = temp / 'source'
    agent.SOURCE_ROOT.mkdir()
    fixture = Path(__file__).resolve().parents[1] / 'testdata' / '测试发票公开版'
    pdfs = sorted(fixture.glob('*.pdf'))[:2]
    assert len(pdfs) == 2
    for path in pdfs:
        shutil.copyfile(path, agent.SOURCE_ROOT / path.name)
    shutil.copyfile(pdfs[0], agent.SOURCE_ROOT / 'z_duplicate.pdf')
    originals = {p.name: p.read_bytes() for p in agent.SOURCE_ROOT.iterdir()}
    assert not agent.safe_source(temp / 'outside.pdf')
    assert agent.parse_plan('整理上个月的发票', datetime(2026, 1, 4))['month'] == '2025-12'
    assert agent.parse_plan('整理2026年9月包含“测试”的发票')['keywords'] == ['测试']
    with TestClient(app) as client:
        assert client.post('/api/v1/agent/plan', json={'prompt': '整理金额大于100元的发票'}).status_code == 422
        assert client.post('/api/v1/agent/plan', json={'prompt': '整理2026年13月的发票'}).status_code == 422
        assert client.post('/api/v1/agent/plan', headers={'origin': 'https://evil.example'}, json={'prompt': '整理全部发票'}).status_code == 403
        response = client.post('/api/v1/agent/tasks', json={'source_dir': str(agent.SOURCE_ROOT), 'prompt': '整理素材里的全部发票，给我过目，然后导出'})
        assert response.status_code == 200, response.text
        tid = response.json()['id']
        result = client.get(f'/api/v1/agent/tasks/{tid}').json()
        assert result['status'] == 'review', result
        rows = [r for r in result['records'] if r['kind'] == 'matched']
        assert len(rows) == 2, result
        assert sum(r['kind'] == 'duplicate' for r in result['records']) == 1
        assert client.get('/api/v1/agent/tasks/latest').json()['id'] == tid
        with SessionLocal() as db:
            assert db.query(Invoice).count() == 2
            assert db.query(LedgerEntry).count() == 0
        one = rows[0]
        iid = one['invoice']['id']
        # 修改后，旧预览不能直接确认。
        patch = client.patch(f'/api/v1/invoices/{iid}', json={'expense_category': '办公费'})
        assert patch.status_code == 200, patch.text
        old = {'invoice_ids': [iid], 'revisions': {iid: one['revision']}}
        assert client.post(f'/api/v1/agent/tasks/{tid}/confirm', json=old).status_code == 409
        fresh = client.get(f'/api/v1/agent/tasks/{tid}').json()
        row = next(r for r in fresh['records'] if r.get('invoice', {}).get('id') == iid)
        confirmed = client.post(f'/api/v1/agent/tasks/{tid}/confirm', json={'invoice_ids': [iid], 'revisions': {iid: row['revision']}})
        assert confirmed.status_code == 200, confirmed.text
        with SessionLocal() as db:
            assert db.query(LedgerEntry).count() == 1
            assert db.query(Invoice).filter(Invoice.status == 'pending_review').count() == 1
        export = client.post('/api/v1/exports/excel', json={'invoice_ids': [iid], 'only_confirmed': True})
        assert export.status_code == 200 and export.content[:2] == b'PK', export.text[:100]
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export.content))
        cells = ' '.join(str(cell.value) for sheet in book for line in sheet for cell in line)
        assert one['invoice']['invoice_number'] in cells
        other = next(r['invoice']['invoice_number'] for r in rows if r['invoice']['id'] != iid)
        assert other not in cells
        second = client.post('/api/v1/agent/tasks', json={'source_dir': str(agent.SOURCE_ROOT), 'prompt': '整理2025年1月的发票'}).json()
        excluded = client.get(f"/api/v1/agent/tasks/{second['id']}").json()
        assert not any(r['kind'] == 'matched' for r in excluded['records'])
        assert client.post(f"/api/v1/agent/tasks/{second['id']}/confirm", json={'invoice_ids': [iid], 'revisions': {iid: row['revision']}}).status_code == 409
        with SessionLocal() as db:
            assert db.query(Invoice).count() == 2, '重复运行不应新建相同文件的票据'
        agent.MAX_FILES = 1
        limited = client.post('/api/v1/agent/tasks', json={'source_dir': str(agent.SOURCE_ROOT), 'prompt': '整理全部发票'}).json()
        limited = client.get(f"/api/v1/agent/tasks/{limited['id']}").json()
        assert limited['limited'] and limited['total'] == 1
        assert originals == {p.name: p.read_bytes() for p in agent.SOURCE_ROOT.iterdir()}
        # DeepSeek 接口使用模拟官方响应，无网络调用、无真实密钥。
        from app.services import deepseek
        import httpx
        assert client.post('/api/v1/agent/ai/chat', json={'message': '帮我整理发票'}).status_code == 409
        saved = client.post('/api/v1/agent/ai/config', json={'api_key': 'sk-test-not-a-real-key', 'model': 'deepseek-flash', 'enabled': True})
        assert saved.status_code == 200 and 'sk-test' not in saved.text
        assert client.post('/api/v1/agent/ai/config', json={'model': 'deepseek-flash', 'enabled': True}).status_code == 200
        assert deepseek.read_config()['api_key'] == 'sk-test-not-a-real-key'
        draft = {'source_dir': str(agent.SOURCE_ROOT), 'summary': '整理测试发票', 'month': None, 'keywords': [], 'buyer': None, 'seller': None, 'exclude_confirmed': True, 'min_amount': 1, 'max_amount': None}
        def completion(message):
            return httpx.Response(200, json={'choices': [{'message': message, 'finish_reason': 'stop'}]})
        tool_message = {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'call-1', 'type': 'function', 'function': {'name': 'prepare_invoice_task', 'arguments': json.dumps(draft)}}]}
        with mock_remote(deepseek, [completion(tool_message), completion({'role': 'assistant', 'content': '已准备计划，请核对后开始。'})]) as remote:
            chat = client.post('/api/v1/agent/ai/chat', json={'message': '排除已入账票据，整理素材中的其他发票', 'task_id': tid})
            assert chat.status_code == 200, chat.text
            answer = chat.json()
            assert answer['draft']['exclude_confirmed'] and len(answer['trace']) == 1
            upstream = remote.call_args_list[1].kwargs['json']
            assert any(m['role'] == 'tool' for m in upstream['messages'])
            assert 'sk-test' not in json.dumps(upstream)
        inspect_message = {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'call-2', 'type': 'function', 'function': {'name': 'inspect_current_task', 'arguments': '{}'}}]}
        with mock_remote(deepseek, [completion(inspect_message), completion({'role': 'assistant', 'content': '我已查看本地实际处理进度。'})]) as remote:
            follow = client.post('/api/v1/agent/ai/chat', json={'message': '现在处理得怎样？', 'conversation_id': answer['conversation_id'], 'task_id': tid})
            assert follow.status_code == 200, follow.text
            upstream = remote.call_args_list[1].kwargs['json']
            serialized = json.dumps(upstream)
            assert '排除已入账' in json.dumps(upstream, ensure_ascii=False)
            assert iid not in serialized and one['invoice']['invoice_number'] not in serialized
        with mock_remote(deepseek, [httpx.Response(401, text='sk-test-not-a-real-key')]):
            probe = client.post('/api/v1/agent/ai/test')
            assert probe.status_code == 502 and 'sk-test' not in probe.text
        agent.MAX_FILES = 50
        started = client.post(f"/api/v1/agent/ai/conversations/{answer['conversation_id']}/start")
        assert started.status_code == 200, started.text
        ai_task_id = started.json()['id']
        ai_task = client.get(f'/api/v1/agent/tasks/{ai_task_id}').json()
        pending_rows = [r for r in ai_task['records'] if r['kind'] == 'matched']
        assert len(pending_rows) == 1, ai_task
        ai_row = pending_rows[0]
        ai_id = ai_row['invoice']['id']
        done = client.post(f'/api/v1/agent/tasks/{ai_task_id}/confirm', json={'invoice_ids': [ai_id], 'revisions': {ai_id: ai_row['revision']}})
        assert done.status_code == 200, done.text
        assert any(r['kind'] == 'matched' and r.get('invoice', {}).get('id') == ai_id for r in done.json()['records']), '本次确认的票据仍需可见，以便下载失败后重试导出'
        from app.routers.agent_chat import Draft
        from app.routers import agent_chat
        assert client.post(f'/api/v1/invoices/{iid}/revoke').status_code == 200
        def model_export(messages, tools=None):
            last = messages[-1]
            if last['role'] == 'user':
                return {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'query-real', 'type': 'function', 'function': {'name': 'query_ledger', 'arguments': '{}'}}]}
            result = json.loads(last['content'])
            if 'query_id' in result and 'url' not in result:
                return {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'export-real', 'type': 'function', 'function': {'name': 'export_ledger', 'arguments': json.dumps({'query_id': result['query_id']})}}]}
            return {'role': 'assistant', 'content': '账本已经导出，包含待复核记录。'}
        with mock_patch.object(agent_chat.deepseek, 'complete', side_effect=model_export):
            exported = client.post('/api/v1/agent/ai/chat', json={'message': '导出当前账本'})
        assert exported.status_code == 200, exported.text
        artifacts = exported.json()['artifacts']
        preview, download = artifacts
        restored = client.get('/api/v1/agent/ai/conversations/' + exported.json()['conversation_id'])
        assert restored.status_code == 200, restored.text
        restored = restored.json()
        assert any(m['role'] == '你' and m['text'] == '导出当前账本' for m in restored['messages'])
        assert restored['artifacts'][0]['rows'] == preview['rows']
        assert restored['artifacts'][1]['url'] == download['url']
        assert 'reasoning_content' not in json.dumps(restored)
        assert client.get('/api/v1/agent/ai/conversations/latest').json()['conversation_id'] == exported.json()['conversation_id']
        assert client.get('/api/v1/agent/ai/conversations/nonexistent').status_code == 404
        assert preview['count'] == 2 and preview['pending'] == 1 and preview['confirmed'] == 1
        file = client.get(download['url'])
        assert file.status_code == 200 and file.content[:2] == b'PK'
        assert client.get('/api/v1/agent/ai/downloads/not-a-real-id').status_code == 404
        book = load_workbook(BytesIO(file.content))
        cells = ' '.join(str(cell.value) for sheet in book for line in sheet for cell in line)
        assert all(r['invoice']['invoice_number'] in cells for r in rows)
        # 查询后票据变化不能静默导出；跨对话快照不允许使用。
        from app.services import agent_ledger
        from fastapi import HTTPException
        with SessionLocal() as db:
            try:
                agent_ledger.export(db, 'different-conversation', {'query_id': preview['id']})
                raise AssertionError('cross-conversation snapshot accepted')
            except HTTPException as exc:
                assert exc.status_code == 404
        client.patch(f'/api/v1/invoices/{iid}', json={'expense_category': '测试更新分类'})
        with SessionLocal() as db:
            try:
                agent_ledger.export(db, exported.json()['conversation_id'], {'query_id': preview['id']})
                raise AssertionError('stale snapshot accepted')
            except HTTPException as exc:
                assert exc.status_code == 409
            empty, card = agent_ledger.query(db, exported.json()['conversation_id'], {'month': '2001-01'}, None)
            assert empty['count'] == 0
            try:
                agent_ledger.export(db, exported.json()['conversation_id'], {'query_id': empty['query_id']})
                raise AssertionError('empty export accepted')
            except HTTPException as exc:
                assert exc.status_code == 404
        direct = {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'start-direct', 'type': 'function', 'function': {'name': 'start_local_task', 'arguments': json.dumps(draft)}}]}
        with mock_patch.object(agent_chat.deepseek, 'complete', side_effect=[direct, HTTPException(504, 'mock timeout')]):
            executed = client.post('/api/v1/agent/ai/chat', json={'message': '开始整理素材里的发票，排除已入账'})
        assert executed.status_code == 200, executed.text
        assert executed.json()['started_task_id']
        actual = client.get('/api/v1/agent/tasks/' + executed.json()['started_task_id']).json()
        assert actual['status'] == 'review', actual
        assert '失败' in executed.json()['reply']
        print('PASS: conversational ledger query/export/download, exact workbook scope, mixed statuses, stale and cross-conversation guards, empty result, direct task start surviving model timeout')
        # 新日期范围：全年、非连续月份、跨年、闰年与包含边界。
        from app.services.task_scope import TaskScope
        from types import SimpleNamespace
        full_year = agent.parse_plan('整理2025年每个月的发票')
        assert full_year['date_from'] == '2025-01-01' and full_year['date_to'] == '2025-12-31'
        selected_months = agent.parse_plan('整理2025年1月、3月的发票')
        assert selected_months['months'] == ['2025-01', '2025-03']
        crossing = agent.parse_plan('整理2025年12月至2026年2月的发票')
        assert crossing['date_from'] == '2025-12-01' and crossing['date_to'] == '2026-02-28'
        leap = agent.parse_plan('整理2024年1月至2月的发票')
        assert leap['date_to'] == '2024-02-29'
        exact = agent.parse_plan('整理2025-01-15至2025-03-02的发票')
        sample = SimpleNamespace(invoice_number='123', seller_name='测试', buyer_name='测试', invoice_remark='', items=[], invoice_date='2025-01-15')
        assert agent.matches(sample, exact)[0] == 'matched'
        sample.invoice_date = '2025-03-02'
        assert agent.matches(sample, exact)[0] == 'matched'
        sample.invoice_date = '2025-03-03'
        assert agent.matches(sample, exact)[0] == 'excluded'
        sample.invoice_date = '2025-02-10'
        assert agent.matches(sample, selected_months)[0] == 'excluded'
        sample.invoice_date = '2025-03-10'
        assert agent.matches(sample, selected_months)[0] == 'matched'
        for invalid in [{'month': '2025-01', 'months': ['2025-02']}, {'date_from': '2025-03-01', 'date_to': '2025-01-01'}, {'months': ['2025-13']}]:
            try:
                TaskScope(**invalid)
                raise AssertionError('invalid time scope accepted')
            except ValueError:
                pass
        assert client.post('/api/v1/agent/tasks', json={'prompt': '整理全部发票'}).status_code == 422
        other_root = temp / 'separate-invoices'
        other_root.mkdir()
        shutil.copyfile(pdfs[0], other_root / 'only-this.pdf')
        assert agent_chat.checked_directory({'source_dir': str(other_root)}, [f'整理“{other_root}”里的票据']) == str(other_root.resolve())
        try:
            agent_chat.checked_directory({'source_dir': str(temp)}, [f'整理“{other_root}”里的票据'])
            raise AssertionError('broader parent folder accepted')
        except HTTPException as exc:
            assert exc.status_code == 422
        external = client.post('/api/v1/agent/tasks', json={'prompt': '整理2025年的发票', 'source_dir': str(other_root)})
        assert external.status_code == 200, external.text
        external = client.get('/api/v1/agent/tasks/' + external.json()['id']).json()
        assert external['total'] == 1 and external['records'][0]['path'] == 'only-this.pdf'
        assert external['plan']['root'] == str(other_root.resolve())
        print('PASS: full year, disjoint months, cross-year, leap-day/endpoints, missing path rejection, explicit alternate directory, no parent-scope expansion')
        # 仅出现在会计科目的词也应被普通搜索、AI查询及导出匹配。
        marker = '唯一测试科目防撞'
        assert client.patch(f'/api/v1/invoices/{iid}', json={'account_subject': marker}).status_code == 200
        found = client.get('/api/v1/ledger', params={'q': marker}).json()
        assert found['total'] == 1 and found['items'][0]['invoice_id'] == iid, found
        all_rows = client.get('/api/v1/ledger').json()['items']
        for index, row in enumerate(all_rows):
            location = client.get('/api/v1/ledger/locate/' + row['invoice_id'], params={'page_size': 1})
            assert location.status_code == 200 and location.json()['page'] == index + 1
        assert client.get('/api/v1/ledger/locate/deleted-invoice').status_code == 404
        with SessionLocal() as db:
            ai_found, _ = agent_ledger.query(db, exported.json()['conversation_id'], {'q': marker}, None)
            assert ai_found['count'] == 1
        filtered_file = client.post('/api/v1/exports/excel', json={'q': marker})
        assert filtered_file.status_code == 200
        book = load_workbook(BytesIO(filtered_file.content))
        cells = ' '.join(str(cell.value) for sheet in book for line in sheet for cell in line)
        assert marker in cells and other not in cells
        from app.models import InvoiceItem
        with SessionLocal() as db:
            for n in range(2):
                db.add(InvoiceItem(invoice_id=iid, item_name='多明细搜索测试', amount=0, sort_order=90+n))
            db.commit()
        assert client.get('/api/v1/ledger', params={'q': '多明细搜索测试'}).json()['total'] == 1
        with mock_remote(deepseek, [httpx.ConnectTimeout('no secret logging')]):
            failed_connect = client.post('/api/v1/agent/ai/test')
            assert failed_connect.status_code == 504 and '连接' in failed_connect.text
        with mock_remote(deepseek, [httpx.ReadTimeout('no secret logging')]):
            failed_read = client.post('/api/v1/agent/ai/test')
            assert failed_read.status_code == 504 and '120' in failed_read.text
        log = (temp / 'data' / 'logs' / 'deepseek.log').read_text(encoding='utf-8')
        assert 'phase=connect' in log and 'phase=ReadTimeout' in log
        assert 'sk-test' not in log and marker not in log
        print('PASS: account-subject query/export consistency, item search without duplicate totals, timeout phase diagnostics without secret/content logging')
        try:
            Draft.model_validate({'summary': 'x', 'min_amount': 10, 'max_amount': 1})
            raise AssertionError('Invalid range must be rejected')
        except ValueError:
            pass
        print('PASS: DeepSeek configuration redaction, tool-call loop, multi-turn history, local summary minimization, upstream error redaction, validated filters')
    print('PASS: local search, real PDF extraction, scope filtering, duplicate reuse, persisted task, stale-review rejection, selected confirmation, scoped Excel export, scan limit, source preservation, origin guard')


try:
    main()
finally:
    engine.dispose()
    from app.services.deepseek import logger
    for handler in list(logger.handlers):
        handler.close()
        logger.removeHandler(handler)
    scratch.cleanup()
